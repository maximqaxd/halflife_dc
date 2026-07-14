#!/usr/bin/env python3
"""Build <file>_checklist.xlsx in the dc_draw checklist format: every function of a
src file present in the binary, with obj/binary match %, whether it's exact (V) or
not (X), plus a "Removed (not in binary)" sheet of source-only functions and rows
for binary functions still MISSING from src.

Usage: python build_checklist_xlsx.py common 0x37eb0 0x3b524 engine/common.c
"""
import os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
name, lo, hi, srcrel = sys.argv[1], int(sys.argv[2], 16), int(sys.argv[3], 16), sys.argv[4]
SYMS = os.path.join(HERE, "symbols.tsv")
DIFF = os.path.join(HERE, "portdata", "%s_objdiff.txt" % name)
OUT = os.path.join(HERE, "portdata", "%s_checklist.xlsx" % name)
SRC = os.path.join(REPO, "src", srcrel)

# --- binary symbols (name -> addr,size) + region + full name set ---
addr, size, allnames = {}, {}, set()
for ln in open(SYMS, encoding="latin-1"):
    p = ln.rstrip("\n").split("\t")
    if len(p) == 3:
        try: a, s = int(p[1], 16), int(p[2])
        except ValueError: continue
        nm = p[0][6:] if p[0].startswith("maybe_") else p[0]
        allnames.add(nm); addr.setdefault(nm, p[1]); size.setdefault(nm, s)
region = {n for n in addr if lo <= int(addr[n], 16) < hi}

# --- objdiff scores: name -> (obj, bin, exact, struct) ---
scored = {}
txt = open(DIFF, encoding="latin-1", errors="replace").read()
for m in re.finditer(r"=== ([A-Za-z0-9_:~]+) ===\s+obj=(\d+) insns\s+bin=(\d+) insns\s+match=([0-9.]+)%.*?struct=([0-9.]+)%", txt):
    scored[m.group(1)] = (int(m.group(2)), int(m.group(3)), float(m.group(4)), float(m.group(5)))

# --- src function names (crude top-level def scan) for the Removed sheet ---
HDR = re.compile(r"([A-Za-z_]\w*)\s*\([^;{}]*\)\s*(?:const)?\s*$")
KW = ("if", "for", "while", "switch", "return", "sizeof", "else", "do")
def src_funcs(path):
    try: t = open(path, encoding="latin-1").read()
    except OSError: return set()
    t = re.sub(r'"(\\.|[^"\\])*"', '""', re.sub(r"//[^\n]*|/\*.*?\*/", " ", t, flags=re.S))
    out, depth, i, n = set(), 0, 0, len(t)
    while i < n:
        ch = t[i]
        if ch == "{":
            if depth == 0:
                j = i - 1
                while j > 0 and t[j] not in ";}{": j -= 1
                mm = HDR.search(t[j+1:i].strip())
                if mm and mm.group(1) not in KW: out.add(mm.group(1))
            depth += 1
        elif ch == "}":
            depth = max(0, depth - 1)
        i += 1
    return out
sfuncs = src_funcs(SRC)
removed = sorted(f for f in sfuncs if f not in allnames)

# --- rows: union of region + scored (dedup) ---
rows = []
for n in sorted(region | set(scored), key=lambda k: int(addr.get(k, "ffffffff"), 16)):
    if n.startswith(("FUN_", "SUB_", "LAB_")): continue
    if n in scored:
        obj, binn, ex, st = scored[n]
        rows.append([n, addr.get(n, ""), binn, obj, ex, st])
    elif n in region:                    # in binary, not compiled/matched -> missing
        rows.append([n, addr.get(n, ""), None, None, None, None])

def statusnote(ex, st):
    if ex is None: return "MISSING - implement from binary"
    if ex >= 99.9: return "exact"
    if st >= 85: return "close - regalloc/order only"
    if st >= 40: return "partial"
    return "diverged"

# sort: exact first (desc), then the rest worst-first
rows.sort(key=lambda r: (0 if (r[4] is not None and r[4] >= 99.9) else 1,
                         -(r[4] or 0) if (r[4] is not None and r[4] >= 99.9) else (r[5] if r[5] is not None else -1)))

# --- styles (match dc_draw) ---
A = "Arial"
hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(name=A, bold=True, color="FFFFFF", size=11)
title_font = Font(name=A, bold=True, size=16, color="1F3864")
sub_font = Font(name=A, italic=True, size=10, color="555555")
v_font = Font(name=A, bold=True, color="0B6E1E", size=11)
x_font = Font(name=A, bold=True, color="9C5700", size=11)
cell_font = Font(name=A, size=10)
thin = Side(style="thin", color="D9D9D9")
border = Border(thin, thin, thin, thin)
def match_fill(m):
    if m is None: return PatternFill("solid", fgColor="FFC7CE")
    if m >= 99.9: return PatternFill("solid", fgColor="A9D08E")
    if m >= 60: return PatternFill("solid", fgColor="C6EFCE")
    if m >= 40: return PatternFill("solid", fgColor="FFEB9C")
    if m >= 20: return PatternFill("solid", fgColor="FCE4D6")
    return PatternFill("solid", fgColor="FFC7CE")

wb = Workbook(); ws = wb.active; ws.title = "%s checklist" % name
ws["A1"] = "Half-Life Dreamcast - %s reconstruction checklist" % os.path.basename(srcrel)
ws["A1"].font = title_font
ws["A2"] = ("V = exact obj==binary   |   X = not yet exact   |   Exact% = objdiff obj-vs-binary   |   "
            "Struct% = same instruction set ignoring register/spill choices (regalloc-invariant)")
ws["A2"].font = sub_font
ws["A3"] = "Binary: RE/HALFLIFE_DC.EXE  (SH-4)   |   Source: src/%s" % srcrel
ws["A3"].font = sub_font

HR = 5
for c, h in enumerate(["#", "Function", "Binary addr", "Bin insns", "Obj insns", "Exact %", "Struct %", "Done", "Notes"], 1):
    cell = ws.cell(HR, c, h); cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = border

r = HR + 1
for i, (nm, ad, binn, obj, ex, st) in enumerate(rows, 1):
    done = "V" if (ex is not None and ex >= 99.9) else "X"
    vals = [i, nm, ("0x" + ad) if ad else "", binn if binn else "", obj if obj else "",
            (ex / 100.0) if ex is not None else "", (st / 100.0) if st is not None else "",
            done, statusnote(ex, st)]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v); cell.font = cell_font; cell.border = border
        if c in (6, 7) and isinstance(v, float): cell.number_format = "0%"; cell.alignment = Alignment(horizontal="center")
        if c == 6: cell.fill = match_fill(ex)
        if c == 8:
            cell.font = v_font if done == "V" else x_font
            cell.fill = PatternFill("solid", fgColor="C6EFCE" if done == "V" else "F2F2F2")
            cell.alignment = Alignment(horizontal="center")
        if c in (1, 3, 4, 5): cell.alignment = Alignment(horizontal="center")
    r += 1
for col, w in zip("ABCDEFGHI", (5, 30, 12, 10, 10, 9, 9, 7, 60)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A6"; ws.auto_filter.ref = "A5:I%d" % (r - 1)

# --- Removed sheet ---
ws2 = wb.create_sheet("Removed (not in binary)")
ws2["A1"] = "Source-only functions (no binary equivalent - candidates to strip)"
ws2["A1"].font = title_font
for c, h in enumerate(["Function", "Note"], 1):
    cell = ws2.cell(3, c, h); cell.fill = hdr_fill; cell.font = hdr_font; cell.border = border
for i, f in enumerate(removed, 4):
    ws2.cell(i, 1, f).font = cell_font; ws2.cell(i, 2, "in src, not in binary symbols").font = cell_font
ws2.column_dimensions["A"].width = 32; ws2.column_dimensions["B"].width = 40

# counts summary in title area
nex = sum(1 for _, _, _, _, ex, _ in rows if ex is not None and ex >= 99.9)
nmiss = sum(1 for _, _, b, _, ex, _ in rows if ex is None)
ws["G1"] = "%d/%d exact, %d missing, %d src-only" % (nex, len(rows), nmiss, len(removed))
ws["G1"].font = sub_font
wb.save(OUT)
print("wrote %s : %d fns (%d exact, %d missing), %d src-only removed" % (OUT, len(rows), nex, nmiss, len(removed)))
